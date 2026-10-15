import importlib.util
import os
import sqlite3
import tempfile
import unittest
from datetime import datetime, timezone, timedelta
from types import SimpleNamespace
from unittest.mock import AsyncMock
from aiogram.fsm.storage.base import StorageKey
from aiogram import Bot, Dispatcher
from aiogram.client.session.base import BaseSession
from aiogram.types import Chat, Message, Update, User, CallbackQuery
from pathlib import Path


os.environ.setdefault("BOT_TOKEN", "123456:abcdefghijklmnopqrstuvwxyzABCDEFGHI")
_tmp = tempfile.TemporaryDirectory()
os.environ["DATABASE_PATH"] = str(Path(_tmp.name) / "test.sqlite3")

spec = importlib.util.spec_from_file_location("remarenko_free_bot", Path(__file__).with_name("bot.py"))
bot = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(bot)


class RecommendationTests(unittest.TestCase):
    def test_beginner_self_study_gets_basic(self):
        tariff, _ = bot.recommend_tariff("new", "200", "start", "self")
        self.assertEqual(tariff, "BASIC")

    def test_first_deals_and_group_support_gets_group(self):
        tariff, _ = bot.recommend_tariff("first", "500", "system", "group")
        self.assertEqual(tariff, "GROUP")

    def test_active_scaling_with_personal_support_gets_master(self):
        tariff, _ = bot.recommend_tariff("active", "more", "scale", "personal")
        self.assertEqual(tariff, "MASTER")


class PersistenceTests(unittest.TestCase):
    def setUp(self):
        bot.connect().close()
        with sqlite3.connect(bot.DATABASE_PATH) as db:
            for table in ("scheduled_messages", "funnel_events", "settings", "users", "form_state", "lead_notifications", "channel_joins"):
                db.execute(f"DELETE FROM {table}")
            db.commit()

    def test_schema_and_timer_schedule_persist(self):
        with bot.connect() as db:
            db.execute(
                """
                INSERT INTO users
                    (telegram_id, source, stage, subscribed, created_at, updated_at)
                VALUES (?, 'test', 'started', 0, ?, ?)
                """,
                (1, bot.now_iso(), bot.now_iso()),
            )
            db.commit()
        bot.schedule_funnel(1)
        with bot.connect() as db:
            rows = db.execute(
                "SELECT kind FROM scheduled_messages WHERE user_id=1 ORDER BY due_at"
            ).fetchall()
        self.assertEqual(
            [row["kind"] for row in rows],
            [f"stage:1:welcome:{i}" for i in range(1,6)],
        )

    def test_test_mode_resets_all_five_timers(self):
        with bot.connect() as db:
            db.execute(
                """
                INSERT INTO users
                    (telegram_id, source, stage, subscribed, created_at, updated_at)
                VALUES (?, 'test', 'started', 0, ?, ?)
                """,
                (2, bot.now_iso(), bot.now_iso()),
            )
            db.commit()
        bot.schedule_funnel(2)
        bot.schedule_funnel(2, test_mode=True)
        with bot.connect() as db:
            count = db.execute(
                """
                SELECT COUNT(*) FROM scheduled_messages
                WHERE user_id=2 AND sent_at IS NULL AND canceled_at IS NULL
                """
            ).fetchone()[0]
        self.assertEqual(count, 5)

    def test_reopening_lesson_does_not_regress_progress(self):
        with bot.connect() as db:
            db.execute("INSERT INTO users(telegram_id,stage,created_at,updated_at) VALUES(7,'recommendation_shown',?,?)", (bot.now_iso(),bot.now_iso()))
        bot.update_user(7, stage='lesson_video_sent', subscribed=1)
        self.assertEqual(bot.get_user(7)['stage'], 'recommendation_shown')
        self.assertTrue(bot.should_skip_scheduled('tariff_offer', bot.get_user(7)))
        self.assertTrue(bot.should_skip_scheduled('lesson_reminder', bot.get_user(7)))

    def test_application_is_stored_with_durable_notification(self):
        with bot.connect() as db:
            db.execute("INSERT INTO users(telegram_id,created_at,updated_at) VALUES(8,?,?)", (bot.now_iso(),bot.now_iso()))
        bot.schedule_funnel(8)
        bot.store_application(8, '@student', 'Lead text')
        self.assertEqual(bot.get_user(8)['lead_status'], 'new')
        with bot.connect() as db:
            self.assertEqual(db.execute('SELECT text FROM lead_notifications WHERE user_id=8').fetchone()[0], 'Lead text')
            self.assertEqual(db.execute('SELECT COUNT(*) FROM scheduled_messages WHERE user_id=8 AND canceled_at IS NULL').fetchone()[0], 0)

    def test_five_form_reminders_do_not_restart_on_repeat_click(self):
        with bot.connect() as db:
            db.execute('INSERT INTO users(telegram_id,stage,created_at,updated_at) VALUES(90,?,?,?)',('lesson_completed',bot.now_iso(),bot.now_iso()))
        bot.schedule_form_reminder(90)
        with bot.connect() as db:
            before=[tuple(r) for r in db.execute('SELECT kind,due_at,created_at FROM scheduled_messages WHERE user_id=90 ORDER BY due_at')]
        bot.schedule_form_reminder(90)
        with bot.connect() as db:
            after=[tuple(r) for r in db.execute('SELECT kind,due_at,created_at FROM scheduled_messages WHERE user_id=90 ORDER BY due_at')]
        self.assertEqual(before,after)
        self.assertEqual(len(before),5)
        hours=[round((datetime.fromisoformat(r[1])-datetime.fromisoformat(r[2])).total_seconds()/3600) for r in before]
        self.assertEqual(hours,[1,5,24,48,72])

    def test_overdue_form_reminders_are_not_sent_in_a_burst(self):
        with bot.connect() as db:
            db.execute('INSERT INTO users(telegram_id,stage,created_at,updated_at) VALUES(90,?,?,?)',('lesson_completed',bot.now_iso(),bot.now_iso()))
        bot.schedule_form_reminder(90)
        with bot.connect() as db:
            db.execute('UPDATE scheduled_messages SET due_at=? WHERE user_id=90',((bot.utc_now()-timedelta(hours=1)).isoformat(),))
        bot.collapse_overdue_form_reminders()
        with bot.connect() as db:
            rows=db.execute('SELECT kind FROM scheduled_messages WHERE user_id=90 AND canceled_at IS NULL').fetchall()
        self.assertEqual([r['kind'] for r in rows],['stage:1:form_start:5'])


class AsyncRegressionTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        bot.connect().close()
        with bot.connect() as db:
            for table in ('users','funnel_events','scheduled_messages','lead_notifications','form_state','channel_joins'):
                db.execute(f'DELETE FROM {table}')
            db.execute('INSERT INTO users(telegram_id,created_at,updated_at) VALUES(50,?,?)', (bot.now_iso(),bot.now_iso()))

    def test_search_unicode_literal_and_combined_filters(self):
        with bot.connect() as db:
            db.execute("UPDATE users SET first_name='Михайло',username='Test_100%',recommendation='GROUP',lesson_received_at=?,questionnaire_started_at=? WHERE telegram_id=50",(bot.now_iso(),bot.now_iso()))
        view={**bot.resolve_lead_view('people'),'search':'МИХАЙЛО','tariff':'GROUP','stage':'started'}
        self.assertEqual(bot.lead_rows(view,0)[0],1)
        view['search']='@test_100%'
        self.assertEqual(bot.lead_rows(view,0)[0],1)
        view['search']="' OR 1=1 --"
        self.assertEqual(bot.lead_rows(view,0)[0],0)
        view['search']=''; view['tariff']='BASIC'
        self.assertEqual(bot.lead_rows(view,0)[0],0)

    async def test_cover_stays_first_and_old_watch_button_cannot_skip_questionnaire(self):
        bot.set_setting('lesson_cover_file_id','stored_cover')
        message=SimpleNamespace(chat=SimpleNamespace(id=50),answer_photo=AsyncMock(),answer=AsyncMock())
        await bot.send_lesson_card(message)
        kwargs=message.answer_photo.call_args.kwargs
        self.assertEqual(kwargs['photo'],'stored_cover')
        self.assertEqual(kwargs['reply_markup'].inline_keyboard[0][0].callback_data,'resume_form')
        self.assertIn('4 короткі запитання',kwargs['caption'])
        self.assertFalse(any(button.callback_data=='resume_form' for row in bot.main_menu(50).inline_keyboard for button in row))
        callback=SimpleNamespace(from_user=SimpleNamespace(id=50),answer=AsyncMock(),message=message,bot=SimpleNamespace())
        state=SimpleNamespace(set_state=AsyncMock())
        await bot.watch_lesson(callback,state)
        self.assertIsNone(bot.get_user(50)['lesson_received_at'])
        bot.update_user(50,stage='recommendation_shown',recommendation='BASIC')
        await bot.send_lesson_card(message)
        self.assertEqual(message.answer_photo.call_args.kwargs['reply_markup'].inline_keyboard[0][0].callback_data,'watch_lesson')
        await bot.watch_lesson(callback,state)
        state.set_state.assert_awaited_with(bot.LeadForm.contact)
        self.assertIsNone(bot.get_user(50)['lesson_received_at'])

    async def test_1000_records_page_and_full_excel_export(self):
        import io
        from openpyxl import load_workbook
        with bot.connect() as db:
            db.execute('DELETE FROM users')
            db.executemany('INSERT INTO users(telegram_id,first_name,recommendation,created_at,updated_at) VALUES(?,?,?,?,?)',
                [(i,'=HYPERLINK("bad")' if i==1 else f'Учень {i}','GROUP',bot.now_iso(),bot.now_iso()) for i in range(1,1001)])
        view=bot.resolve_lead_view('all')
        count,page,rows=bot.lead_rows(view,124)
        self.assertEqual((count,page,len(rows)),(1000,124,8))
        message=SimpleNamespace(answer=AsyncMock())
        await bot.show_leads(message,'all',124)
        self.assertIn('Сторінка 125 із 125',message.answer.call_args.args[0])
        _,_,allrows=bot.lead_rows(view)
        self.assertEqual(len(allrows),1000)
        book=load_workbook(io.BytesIO(bot.excel_bytes(allrows)))
        sheet=book.active
        self.assertEqual(sheet.max_row,1001)
        formula_row=next(row for row in sheet.iter_rows(min_row=2) if row[0].value=='1')
        self.assertEqual(formula_row[1].value,'=HYPERLINK("bad")')
        self.assertEqual(formula_row[1].data_type,'s')
        self.assertEqual(sheet.freeze_panes,'A2')
        self.assertEqual(sheet.auto_filter.ref,'A1:N1001')

    async def test_filtered_context_survives_card_status_and_back(self):
        bot.update_user(50,recommendation='GROUP')
        token=bot.save_lead_view({**bot.resolve_lead_view('people'),'tariff':'GROUP'})
        self.assertEqual(bot.resolve_lead_view(token)['tariff'],'GROUP')
        msg=SimpleNamespace(edit_text=AsyncMock())
        callback=SimpleNamespace(from_user=SimpleNamespace(id=bot.ADMIN_CHAT_ID),answer=AsyncMock(),message=msg,data=f'adminlead:detail:50:{token}:0')
        await bot.admin_lead_callback(callback)
        markup=msg.edit_text.call_args.kwargs['reply_markup']
        callback.data=markup.inline_keyboard[1][0].callback_data
        self.assertLessEqual(len(callback.data.encode()),64)
        await bot.admin_lead_callback(callback)
        back=msg.edit_text.call_args.kwargs['reply_markup'].inline_keyboard[-1][0].callback_data
        self.assertEqual(back,f'adminlead:list:{token}:0')
        callback.data=back
        await bot.admin_lead_callback(callback)
        self.assertIn('Тариф: Груповий',msg.edit_text.call_args.args[0])

    async def test_owner_search_input_and_export_permissions(self):
        message=SimpleNamespace(answer=AsyncMock(),answer_document=AsyncMock(),edit_text=AsyncMock())
        callback=SimpleNamespace(from_user=SimpleNamespace(id=51),answer=AsyncMock(),message=message,data='av:export:people')
        await bot.admin_view_callback(callback)
        message.answer_document.assert_not_awaited()
        callback.from_user.id=bot.ADMIN_CHAT_ID; callback.data='av:search:people'
        await bot.admin_view_callback(callback)
        message.from_user=SimpleNamespace(id=bot.ADMIN_CHAT_ID); message.text='50'
        self.assertTrue(bot.is_admin_list_input(message))
        await bot.admin_list_input(message)
        self.assertIn('Пошук: 50',message.answer.call_args.args[0])
        self.assertFalse(bot.is_admin_list_input(message))
        callback.data='av:export:'+bot.get_setting('admin_current_view')
        await bot.admin_view_callback(callback)
        document=message.answer_document.call_args.args[0]
        self.assertTrue(document.filename.endswith('.xlsx'))

    def test_today_uses_kyiv_midnight_and_original_entry(self):
        from unittest.mock import patch
        fixed=datetime(2026,10,15,0,30,tzinfo=timezone.utc)
        with bot.connect() as db:
            db.execute('UPDATE users SET created_at=?,updated_at=? WHERE telegram_id=50',('2026-10-14T22:00:00+00:00',fixed.isoformat()))
        view={**bot.resolve_lead_view('people'),'date':'today'}
        with patch.object(bot,'utc_now',return_value=fixed):
            self.assertEqual(bot.lead_rows(view,0)[0],1)
        with bot.connect() as db:
            db.execute('UPDATE users SET created_at=? WHERE telegram_id=50',('2026-10-14T20:00:00+00:00',))
        with patch.object(bot,'utc_now',return_value=fixed):
            self.assertEqual(bot.lead_rows(view,0)[0],0)

    def test_new_step_restarts_five_deadlines_and_old_series_is_stale(self):
        from unittest.mock import patch
        anchor=datetime(2026,10,15,10,tzinfo=timezone.utc)
        with patch.object(bot,'utc_now',return_value=anchor):
            bot.update_user(50,stage='questionnaire_started')
        original=bot.get_user(50)
        self.assertEqual(original['reminder_step'],'form_exp')
        old_kind=f'stage:{original["reminder_generation"]}:form_exp:1'
        with patch.object(bot,'utc_now',return_value=anchor+timedelta(hours=3)):
            bot.update_user(50,experience='Новачок')
        user=bot.get_user(50)
        self.assertEqual(user['reminder_step'],'form_cap')
        self.assertTrue(bot.should_skip_scheduled(old_kind,user))
        with bot.connect() as db:
            rows=db.execute('SELECT * FROM scheduled_messages WHERE user_id=50 AND canceled_at IS NULL ORDER BY due_at').fetchall()
        self.assertEqual(len(rows),5)
        self.assertEqual([int((datetime.fromisoformat(r['due_at'])-anchor-timedelta(hours=3)).total_seconds()/3600) for r in rows],[1,5,24,48,72])
        before=[(r['kind'],r['due_at']) for r in rows]
        bot.update_user(50,lead_status='work')
        bot.schedule_funnel(50)
        bot.initialize_stage_reminders()
        bot.initialize_stage_reminders()
        with bot.connect() as db:
            after=[tuple(r) for r in db.execute('SELECT kind,due_at FROM scheduled_messages WHERE user_id=50 AND canceled_at IS NULL ORDER BY due_at')]
        self.assertEqual(before,after)

    def test_all_funnel_steps_continue_from_contact_to_lesson(self):
        stages=[({'stage':'subscription_required'},'subscribe'),
                ({'stage':'lesson_opened','subscribed':1},'form_start'),
                ({'stage':'lesson_video_sent'},'watch'),
                ({'stage':'lesson_completed'},'form_start'),
                ({'stage':'questionnaire_started'},'form_exp'),
                ({'experience':'Новачок'},'form_cap'),
                ({'capital':'До $200'},'form_goal'),
                ({'goal':'Стартувати з нуля'},'form_support'),
                ({'support':'Навчатися самостійно','recommendation':'BASIC','stage':'recommendation_shown'},'contact'),
                ({'stage':'contact_requested'},'contact')]
        for values,expected in stages:
            bot.update_user(50,**values)
            user=bot.get_user(50)
            self.assertEqual(user['reminder_step'],expected)
            with bot.connect() as db:
                rows=db.execute('SELECT kind FROM scheduled_messages WHERE user_id=50 AND canceled_at IS NULL AND sent_at IS NULL').fetchall()
            self.assertEqual(len(rows),5)
            self.assertTrue(all(expected in r['kind'] for r in rows))
            self.assertEqual(len(set(bot.REMINDER_COPY[expected])),5)
        bot.store_application(50,'@student','Application')
        self.assertEqual(bot.get_user(50)['reminder_step'],'lesson_ready')
        with bot.connect() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM scheduled_messages WHERE user_id=50 AND canceled_at IS NULL AND sent_at IS NULL').fetchone()[0],5)
        bot.record_event(50,'lesson_video_sent')
        bot.schedule_funnel(50)
        self.assertEqual(bot.get_user(50)['reminder_step'],'watch')
        bot.record_event(50,'lesson_completed')
        bot.schedule_funnel(50)
        self.assertEqual(bot.get_user(50)['reminder_step'],'completed')
        with bot.connect() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM scheduled_messages WHERE user_id=50 AND canceled_at IS NULL AND sent_at IS NULL').fetchone()[0],0)

    async def test_retired_button_restores_disabled_users_once_and_old_callback_cannot_stop(self):
        with bot.connect() as db:
            db.execute("DELETE FROM settings WHERE key='reminder_controls_removed_216'")
        bot.update_user(50,reminders_enabled=0)
        bot.update_user(50,stage='questionnaire_started',experience='Новачок')
        bot.initialize_stage_reminders()
        callback=SimpleNamespace(from_user=SimpleNamespace(id=50),answer=AsyncMock(),message=SimpleNamespace(answer=AsyncMock()))
        with bot.connect() as db:
            before=[tuple(r) for r in db.execute('SELECT kind,due_at FROM scheduled_messages WHERE user_id=50 AND canceled_at IS NULL')]
        self.assertEqual(len(before),5)
        self.assertEqual(bot.get_user(50)['reminders_enabled'],1)
        await bot.stop_reminders(callback)
        self.assertEqual(bot.get_user(50)['reminders_enabled'],1)
        self.assertFalse(any(b.callback_data in {'stop_reminders','enable_reminders'} for row in bot.main_menu(50).inline_keyboard for b in row))
        with bot.connect() as db:
            row=db.execute('SELECT kind FROM scheduled_messages WHERE user_id=50 AND canceled_at IS NULL LIMIT 1').fetchone()
        api=SimpleNamespace(send_message=AsyncMock())
        self.assertTrue(await bot.send_scheduled_message(api,row['kind'],50,bot.get_user(50)))
        buttons=[button.callback_data for line in api.send_message.call_args.kwargs['reply_markup'].inline_keyboard for button in line]
        self.assertNotIn('stop_reminders',buttons)
        bot.initialize_stage_reminders()
        with bot.connect() as db:
            after=[tuple(r) for r in db.execute('SELECT kind,due_at FROM scheduled_messages WHERE user_id=50 AND canceled_at IS NULL')]
        self.assertEqual(before,after)

    async def test_stale_queue_item_cannot_send_after_progress(self):
        bot.update_user(50,stage='questionnaire_started')
        stale=bot.get_user(50)
        with bot.connect() as db:
            kind=db.execute('SELECT kind FROM scheduled_messages WHERE user_id=50 AND canceled_at IS NULL LIMIT 1').fetchone()[0]
        bot.update_user(50,experience='Новачок')
        api=SimpleNamespace(send_message=AsyncMock())
        self.assertFalse(await bot.send_scheduled_message(api,kind,50,stale))
        api.send_message.assert_not_awaited()
        current=bot.get_user(50)
        with bot.connect() as db:
            kind=db.execute('SELECT kind FROM scheduled_messages WHERE user_id=50 AND canceled_at IS NULL ORDER BY due_at LIMIT 1').fetchone()[0]
        self.assertTrue(await bot.send_scheduled_message(api,kind,50,current))
        self.assertIn('капітал',api.send_message.call_args.args[1])

    async def test_missing_legacy_fsm_recovers_from_saved_answers(self):
        bot.update_user(50,stage='questionnaire_started',experience='Новачок',capital='До $200')
        state=SimpleNamespace(get_state=AsyncMock(return_value=None),update_data=AsyncMock(),set_state=AsyncMock())
        callback=SimpleNamespace(from_user=SimpleNamespace(id=50),answer=AsyncMock(),message=SimpleNamespace(answer=AsyncMock()))
        await bot.resume_form(callback,state)
        state.set_state.assert_awaited_once_with(bot.LeadForm.goal)
        self.assertEqual(state.update_data.call_args.kwargs['capital_code'],'200')
        self.assertIn('3/4',callback.message.answer.call_args.args[0])

    def test_snapshot_contains_committed_wal_state_and_schedules(self):
        bot.update_user(50,stage='questionnaire_started',experience='Новачок')
        bot.set_setting('lesson_video_file_id','test_video')
        data=bot.database_snapshot()
        restored=sqlite3.connect(':memory:')
        try:
            restored.deserialize(data)
            self.assertEqual(restored.execute('PRAGMA integrity_check').fetchone()[0],'ok')
            self.assertEqual(restored.execute('SELECT experience FROM users WHERE telegram_id=50').fetchone()[0],'Новачок')
            self.assertEqual(restored.execute('SELECT COUNT(*) FROM scheduled_messages WHERE canceled_at IS NULL AND user_id=50').fetchone()[0],5)
        finally:
            restored.close()

    def test_legacy_database_migration_preserves_people_media_answers_and_sent_timers(self):
        previous=bot.DATABASE_PATH
        with tempfile.TemporaryDirectory() as folder:
            bot.DATABASE_PATH=str(Path(folder)/'legacy.sqlite3')
            try:
                with sqlite3.connect(bot.DATABASE_PATH) as db:
                    db.execute('''CREATE TABLE users(telegram_id INTEGER PRIMARY KEY,username TEXT,first_name TEXT,
                        source TEXT NOT NULL DEFAULT 'direct',stage TEXT NOT NULL DEFAULT 'started',subscribed INTEGER NOT NULL DEFAULT 0,
                        experience TEXT,capital TEXT,goal TEXT,contact TEXT,created_at TEXT NOT NULL,updated_at TEXT NOT NULL)''')
                    db.execute("INSERT INTO users(telegram_id,first_name,experience,stage,created_at,updated_at) VALUES(777,'Старий клієнт','Новачок','questionnaire_started',?,?)",(bot.now_iso(),bot.now_iso()))
                bot.connect().close()
                bot.set_setting('lesson_video_file_id','legacy_video')
                bot.initialize_stage_reminders()
                user=bot.get_user(777)
                self.assertEqual((user['first_name'],user['experience'],user['reminder_step']),('Старий клієнт','Новачок','form_cap'))
                with bot.connect() as db:
                    db.execute('UPDATE scheduled_messages SET sent_at=? WHERE user_id=777 AND kind LIKE ?', (bot.now_iso(),'%:1'))
                    before=[tuple(r) for r in db.execute('SELECT kind,due_at,sent_at FROM scheduled_messages ORDER BY id')]
                bot.connect().close()
                bot.initialize_stage_reminders()
                self.assertEqual(bot.get_setting('lesson_video_file_id'),'legacy_video')
                with bot.connect() as db:
                    after=[tuple(r) for r in db.execute('SELECT kind,due_at,sent_at FROM scheduled_messages ORDER BY id')]
                self.assertEqual(before,after)
            finally:
                bot.DATABASE_PATH=previous

    async def test_worker_block_cancels_all_remaining_reminders(self):
        import asyncio
        from unittest.mock import patch
        from aiogram.exceptions import TelegramForbiddenError
        from aiogram.methods import SendMessage
        bot.update_user(50,stage='questionnaire_started')
        with bot.connect() as db:
            db.execute('UPDATE scheduled_messages SET due_at=? WHERE user_id=50',((bot.utc_now()-timedelta(minutes=1)).isoformat(),))
        api=SimpleNamespace(send_message=AsyncMock(side_effect=TelegramForbiddenError(method=SendMessage(chat_id=50,text='test'),message='Forbidden: bot was blocked by the user')))
        with patch.object(bot.asyncio,'sleep',new=AsyncMock(side_effect=asyncio.CancelledError)):
            with self.assertRaises(asyncio.CancelledError):
                await bot.funnel_worker(api)
        self.assertEqual(bot.get_user(50)['reminders_enabled'],0)
        with bot.connect() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM scheduled_messages WHERE user_id=50 AND canceled_at IS NULL').fetchone()[0],0)

    async def test_backup_and_storage_commands_owner_only(self):
        message=SimpleNamespace(from_user=SimpleNamespace(id=51),answer=AsyncMock(),answer_document=AsyncMock())
        await bot.backup_db(message)
        await bot.storage_status(message)
        message.answer.assert_not_awaited()
        message.answer_document.assert_not_awaited()

    async def test_form_survives_new_storage_instance(self):
        key = StorageKey(bot_id=123456,chat_id=50,user_id=50)
        first=bot.SQLiteStorage()
        await first.set_state(key,bot.LeadForm.capital)
        await first.set_data(key,{'experience':'Новачок','experience_code':'new'})
        await first.close()
        second=bot.SQLiteStorage()
        self.assertEqual(await second.get_state(key),bot.LeadForm.capital.state)
        self.assertEqual((await second.get_data(key))['experience_code'],'new')

    async def test_failed_lead_notification_retries_and_then_is_sent(self):
        bot.store_application(50,'@user','Stored application')
        api=SimpleNamespace(send_message=AsyncMock(side_effect=RuntimeError('temporary outage')))
        await bot.deliver_lead_notifications(api)
        with bot.connect() as db:
            row=db.execute('SELECT * FROM lead_notifications WHERE user_id=50').fetchone()
            self.assertIsNone(row['sent_at'])
            self.assertEqual(row['attempts'],1)
            db.execute('UPDATE lead_notifications SET due_at=? WHERE user_id=50',(bot.now_iso(),))
        api.send_message=AsyncMock(return_value=SimpleNamespace(message_id=123))
        await bot.deliver_lead_notifications(api)
        with bot.connect() as db:
            self.assertIsNotNone(db.execute('SELECT sent_at FROM lead_notifications WHERE user_id=50').fetchone()[0])
        api.send_message.assert_awaited_once()

    async def test_fallback_link_records_issued_lesson(self):
        previous=bot.LESSON_URL
        bot.LESSON_URL='https://example.com/lesson'
        try:
            message=SimpleNamespace(chat=SimpleNamespace(id=50),answer=AsyncMock())
            await bot.send_lesson_video(message)
            self.assertIsNone(bot.get_user(50)['lesson_received_at'])
            self.assertIn('Спочатку пройди',message.answer.call_args.args[0])
            bot.update_user(50,recommendation='GROUP',stage='recommendation_shown')
            await bot.send_lesson_video(message)
            self.assertIsNone(bot.get_user(50)['lesson_received_at'])
            self.assertIn('залиш номер',message.answer.call_args.args[0])
            bot.update_user(50,contact='@student',stage='application_completed')
            await bot.send_lesson_video(message)
            self.assertIsNotNone(bot.get_user(50)['lesson_received_at'])
        finally: bot.LESSON_URL=previous

    async def test_early_lesson_done_is_rejected(self):
        callback=SimpleNamespace(from_user=SimpleNamespace(id=50),answer=AsyncMock(),message=SimpleNamespace(answer=AsyncMock()))
        await bot.lesson_done(callback)
        callback.message.answer.assert_not_awaited()

    async def test_resume_after_recommendation_requires_contact_before_lesson(self):
        bot.update_user(50,stage='recommendation_shown',recommendation='GROUP')
        state=SimpleNamespace(get_state=AsyncMock(return_value=None),set_state=AsyncMock())
        callback=SimpleNamespace(from_user=SimpleNamespace(id=50,username='student'),answer=AsyncMock(),message=SimpleNamespace(chat=SimpleNamespace(id=50),answer=AsyncMock()),bot=SimpleNamespace())
        from unittest.mock import patch
        with patch.object(bot,'subscribed',new=AsyncMock(return_value=True)), patch.object(bot,'send_lesson_video',new=AsyncMock()) as send_video:
            await bot.resume_form(callback,state)
            send_video.assert_not_awaited()
            bot.update_user(50,contact='@student',stage='application_completed')
            await bot.resume_form(callback,state)
            send_video.assert_awaited_once()
        state.set_state.assert_awaited_with(bot.LeadForm.contact)
        self.assertTrue(callback.message.answer.await_count)

    async def test_questionnaire_notifies_owner_before_contact(self):
        api=SimpleNamespace(send_message=AsyncMock(return_value=SimpleNamespace(message_id=777)))
        message=SimpleNamespace(chat=SimpleNamespace(id=50),answer=AsyncMock(),bot=api)
        state=SimpleNamespace(set_state=AsyncMock())
        await bot.send_recommendation(message, {'experience_code':'new','capital_code':'200','goal_code':'start','support_code':'self'},state)
        self.assertIsNone(bot.get_user(50)['contact'])
        self.assertEqual(bot.get_user(50)['stage'],'contact_requested')
        state.set_state.assert_awaited_once_with(bot.LeadForm.contact)
        first_markup=message.answer.await_args_list[0].kwargs['reply_markup']
        self.assertEqual(sum(len(row) for row in first_markup.inline_keyboard),5)
        self.assertEqual(
            [button.callback_data for row in first_markup.inline_keyboard for button in row],
            ['watch_lesson','rec_details:BASIC','contact_start','rec_site','rec_home'],
        )
        api.send_message.assert_awaited_once()
        view=SimpleNamespace(answer=AsyncMock())
        await bot.show_leads(view,'new')
        self.assertIn('· 1', view.answer.call_args.args[0])

    async def test_each_of_five_recommendation_buttons_requires_contact(self):
        bot.update_user(50,recommendation='GROUP',stage='contact_requested')
        handlers = [
            ('watch_lesson', bot.watch_lesson),
            ('rec_details:GROUP', bot.recommendation_details),
            ('contact_start', bot.contact_start),
            ('rec_site', bot.recommendation_site),
            ('rec_home', bot.recommendation_home),
        ]
        for data,handler in handlers:
            state=SimpleNamespace(set_state=AsyncMock())
            message=SimpleNamespace(chat=SimpleNamespace(id=50),answer=AsyncMock())
            callback=SimpleNamespace(
                from_user=SimpleNamespace(id=50,username='student'),
                answer=AsyncMock(),message=message,data=data,bot=SimpleNamespace(),
            )
            await handler(callback,state)
            state.set_state.assert_awaited_with(bot.LeadForm.contact)
            self.assertTrue(message.answer.await_count)
            self.assertIsNone(bot.get_user(50)['contact'])

    async def test_contact_updates_existing_message_and_preserves_work_status(self):
        bot.update_user(50,recommendation='GROUP',stage='recommendation_shown')
        api=SimpleNamespace(send_message=AsyncMock(return_value=SimpleNamespace(message_id=888)),edit_message_text=AsyncMock())
        bot.queue_questionnaire(50)
        await bot.deliver_lead_notifications(api)
        bot.update_user(50,lead_status='work')
        state=SimpleNamespace(clear=AsyncMock())
        message=SimpleNamespace(
            from_user=SimpleNamespace(id=50,username='student',full_name='Student'),
            text='Використати мій Telegram',contact=None,answer=AsyncMock(),bot=api,
        )
        await bot.apply_contact(message,state)
        api.send_message.assert_awaited_once()
        api.edit_message_text.assert_awaited_once()
        self.assertEqual(api.edit_message_text.call_args.kwargs['message_id'],888)
        self.assertEqual(bot.get_user(50)['lead_status'],'work')
        self.assertEqual(bot.get_user(50)['contact'],'@student')
        self.assertEqual(bot.get_user(50)['reminder_step'],'lesson_ready')
        second_markup=message.answer.await_args_list[1].kwargs['reply_markup']
        self.assertEqual(sum(len(row) for row in second_markup.inline_keyboard),5)

    async def test_random_text_cannot_unlock_contact_gate(self):
        bot.update_user(50,recommendation='BASIC',stage='contact_requested')
        state=SimpleNamespace(clear=AsyncMock())
        message=SimpleNamespace(
            from_user=SimpleNamespace(id=50,username='student',full_name='Student'),
            text='привіт',contact=None,answer=AsyncMock(),bot=SimpleNamespace(),
        )
        await bot.apply_contact(message,state)
        self.assertIsNone(bot.get_user(50)['contact'])
        state.clear.assert_not_awaited()
        self.assertIn('коректний номер',message.answer.call_args.args[0])

    async def test_non_owner_cannot_open_lead_panel(self):
        callback=SimpleNamespace(from_user=SimpleNamespace(id=50),answer=AsyncMock(),message=SimpleNamespace(answer=AsyncMock()),data='adminlead:detail:50')
        await bot.admin_lead_callback(callback)
        callback.message.answer.assert_not_awaited()

    async def test_back_button_keeps_people_page_after_status_change(self):
        message=SimpleNamespace(edit_text=AsyncMock())
        callback=SimpleNamespace(from_user=SimpleNamespace(id=bot.ADMIN_CHAT_ID),answer=AsyncMock(),message=message,data='adminlead:detail:50:people:2')
        await bot.admin_lead_callback(callback)
        markup=message.edit_text.call_args.kwargs['reply_markup']
        self.assertEqual(markup.inline_keyboard[-1][0].callback_data,'adminlead:list:people:2')
        callback.data=markup.inline_keyboard[1][0].callback_data
        await bot.admin_lead_callback(callback)
        markup=message.edit_text.call_args.kwargs['reply_markup']
        self.assertEqual(markup.inline_keyboard[-1][0].callback_data,'adminlead:list:people:2')
        self.assertEqual(bot.get_user(50)['lead_status'],'work')

    async def test_list_cards_include_filter_and_page(self):
        with bot.connect() as db:
            for uid in range(60,70):
                db.execute('INSERT INTO users(telegram_id,created_at,updated_at) VALUES(?,?,?)',(uid,bot.now_iso(),bot.now_iso()))
        message=SimpleNamespace(answer=AsyncMock())
        await bot.show_leads(message,'people',1)
        markup=message.answer.call_args.kwargs['reply_markup']
        self.assertTrue(markup.inline_keyboard[0][0].callback_data.endswith(':people:1'))

    async def test_join_count_is_unique_and_identifies_invite(self):
        previous=bot.CHANNEL_ID
        bot.CHANNEL_ID='-10042'
        try:
            user=SimpleNamespace(id=50,is_bot=False)
            event=SimpleNamespace(chat=SimpleNamespace(id=-10042,username=None),date=datetime.now(timezone.utc),
                old_chat_member=SimpleNamespace(status='left'),new_chat_member=SimpleNamespace(status='member',user=user),
                invite_link=SimpleNamespace(invite_link=bot.CHANNEL_URL))
            await bot.track_channel_join(event)
            await bot.track_channel_join(event)
            with bot.connect() as db:
                count,via=db.execute('SELECT COUNT(*),SUM(via_bot_link) FROM channel_joins').fetchone()
            self.assertEqual((count,via),(1,1))
        finally: bot.CHANNEL_ID=previous

    async def test_full_questionnaire_through_real_dispatcher(self):
        class FakeSession(BaseSession):
            def __init__(self):
                super().__init__()
                self.sent=[]
            async def close(self): pass
            async def make_request(self, api, method, timeout=None):
                if hasattr(method, 'chat_id') and hasattr(method, 'text'):
                    self.sent.append(method.text)
                    return Message(message_id=len(self.sent),date=datetime.now(timezone.utc),chat=Chat(id=int(method.chat_id),type='private'),text=method.text)
                return True
            async def stream_content(self, *args, **kwargs):
                if False: yield b''
        session=FakeSession()
        api=Bot(os.environ['BOT_TOKEN'],session=session)
        dp=Dispatcher(storage=bot.SQLiteStorage())
        dp.include_router(bot.router)
        uid=100
        async def send_text(text, who=uid):
            msg=Message(message_id=len(session.sent)+1,date=datetime.now(timezone.utc),chat=Chat(id=who,type='private'),from_user=User(id=who,is_bot=False,first_name='Student'),text=text)
            await dp.feed_update(api,Update(update_id=len(session.sent)+1,message=msg))
        async def press(data):
            msg=Message(message_id=1,date=datetime.now(timezone.utc),chat=Chat(id=uid,type='private'),from_user=User(id=123456,is_bot=True,first_name='Bot'),text='menu')
            cb=CallbackQuery(id='query',from_user=User(id=uid,is_bot=False,first_name='Student'),chat_instance='chat',message=msg,data=data)
            await dp.feed_update(api,Update(update_id=len(session.sent)+1,callback_query=cb))
        await send_text('/start instagram')
        for data in ('apply','exp:first','cap:500','goal:system','sup:group'):
            await press(data)
        await send_text('@student')
        user=bot.get_user(uid)
        self.assertEqual(user['stage'],'application_completed')
        self.assertEqual(user['recommendation'],'GROUP')
        self.assertEqual(user['source'],'instagram')
        self.assertEqual(user['contact'],'@student')
        await send_text('/leads',bot.ADMIN_CHAT_ID)
        self.assertTrue(any('Нові' in text for text in session.sent))
        await dp.storage.close()
        await api.session.close()


if __name__ == "__main__":
    unittest.main()
